#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
diamond_table1_dev.py

目的:
  単施設(DIAMOND)コホートのTable1（患者背景表）を作成する。

  ソース: `DIAMOND case file.xlsx` の
    - '本当の除外なし921例' シート（肺瘻0or1ラベル＋PRE-*臨床変数、n=921）を基本母集団とする
    - '臨床因子追加' シート（身長・体重・肺機能等、n=1031）を患者番号で結合し、
      BMI・肺機能変数を追加する

  ⚠️ 重要な前提（2026-08-16、本人確認済み）:
    この921例は、実際に時系列モデルの学習に使われているコホート
    （9:1分割ログでn_total=1020、2026-08-09時点でn=1025）の**部分集合**であり、
    「臨床背景データが揃っている症例のみ」（欠測99〜104例）である。
    したがって本Table1は「モデル学習コホート全体」を100%代表するものではなく、
    **モデル学習コホートのうち背景データ完備例（921/1020、約90%）**という
    limitationを論文Methods/Limitationsに明記する必要がある。
    本スクリプトはこの前提を出力ファイルのヘッダにも明記する。

  '本当の除外なし921例'シートは113列、'臨床因子追加'シートは101列あり、
  過去の列名一覧表示（diamond_inspect_dev_background_candidates.py）では
  60列までしか確認できていない。未確認の残り列に手術時間・出血量・術式等が
  含まれる可能性があるため、本スクリプトは①既知の確定列 ②キーワードでの
  自動検出、の両方を試み、かつ全列名をファイルに書き出して事後確認できるようにする。

  出力は必ずファイル(CSV/Markdown)に保存し、ターミナルには要約のみ出す設計
  （「出力が長すぎて貼れない」問題への対策）。

使い方（推奨: Excelで両シートを事前にCSVエクスポートしてから渡す。xlsx直読みが
  極端に遅い場合があるため）:
  python diamond_table1_dev.py \
      --base_csv ~/Documents/DIAMOND/DIAMOND/table1_source_base.csv \
      --factor_csv ~/Documents/DIAMOND/DIAMOND/table1_source_factor.csv \
      --output_dir ~/Documents/DIAMOND/DIAMOND/split_models_external

使い方（xlsxを直接読む場合、遅い可能性あり）:
  python diamond_table1_dev.py --output_dir ~/Documents/DIAMOND/DIAMOND/split_models_external
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path
from typing import Dict, List, Optional

import numpy as np
import pandas as pd

CASE_FILE_DEFAULT = "~/Documents/DIAMOND/DIAMOND case file.xlsx"
BASE_SHEET = "本当の除外なし921例"
FACTOR_SHEET = "臨床因子追加"

# 既に列名を確認済みの確定候補。
# ⚠️ 2026-08-16実データで判明: 'PRE-S.I.'(喫煙指数)は0/100/200等の生のBrinkman指数
#   数値であり、カテゴリ変数ではなく連続変数として扱う（値ごとに行が爆発するため）。
CONTINUOUS_VARS = [
    ("年齢", "PRE-age"),
    ("喫煙指数(Brinkman Index)", "PRE-S.I."),
]
CATEGORICAL_VARS = [
    ("性別", "PRE-sex"),
    ("喫煙(有無)", "PRE-smoking or never"),
    ("喫煙(現喫煙)", "PRE-smoking"),
    ("COPD", "PRE-COPD"),
    ("間質性肺炎", "PRE-IP"),
    ("糖尿病", "PRE-DM"),
    ("抗血栓薬内服", "PRE-anticoag/plt"),
    ("ステロイド", "PRE-steroid"),
    ("PS", "PRE-PS"),
    ("疾患の種類(1)", "疾患の種類(1)"),
    ("主病巣側", "主病巣側"),
    ("肺瘻(アウトカム, 0/1)", "肺瘻0or1"),
]

# 列名が未確認(60列より後ろ)のため、キーワードで緩く自動検出する候補
# ⚠️ '手術術式'は自由記述に近く(患者ごとにほぼユニーク、大半n=1)、Table1のカテゴリ
#   変数として不適切なため対象から外す。同種の情報は'肺切除範囲'/'切除肺葉'/'Approach'
#   でより粗い（Table1向きの）カテゴリとして既に取得できる。
KEYWORD_CANDIDATES = {
    "continuous": ["手術時間", "出血量", "surgtime", "bloodloss"],
    "categorical": [
        "手術の種類", "Approach", "アプローチ", "肺切除範囲",
        "切除肺葉", "開胸法", "ﾘﾝﾊﾟ節郭清",
    ],
}

FACTOR_CONTINUOUS_VARS = [
    ("身長(cm)", "身長(cm)"),
    ("体重(kg)", "体重(kg)"),
    ("VC", "PRE-VC"),
    ("FVC", "PRE-FVC"),
    ("FVC(%)", "PRE-FVC(%)"),
    ("FEV1.0", "PRE-FEV1.0"),
    ("FEV1.0(%)", "PRE-FEV1.0(%)"),
]
FACTOR_CATEGORICAL_VARS = [
    ("術中Leak test", "INTRA-Leak test(0=leakなし、1=あり/消失、2=あり/消失確認なし)"),
    ("術中sealant", "INTRA-sealant(0=なし、1=のりのみ、2=シートのみ、3=のり＋シート、4=縫縮あり)"),
    ("胸腔内癒着", "INTRA-adhesion"),
]

JOIN_KEY_CANDIDATES = ["患者番号", "通し番号"]

# 参考: 実際にモデル学習に使われている時系列コホートの規模（このファイル単体では検証不可）
MODEL_COHORT_N_REFERENCE = {
    "2026-08-09_frozen_model_base": 1025,
    "2026-08-15_9to1_split_log_base": 1020,
}


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser()
    p.add_argument("--case_file", default=CASE_FILE_DEFAULT, help="xlsxから直接読む場合(遅い場合あり)")
    p.add_argument("--base_csv", default=None, help=f"'{BASE_SHEET}'シートを事前にCSVエクスポートしたファイル(推奨、高速)")
    p.add_argument("--factor_csv", default=None, help=f"'{FACTOR_SHEET}'シートを事前にCSVエクスポートしたファイル(推奨、高速)")
    p.add_argument("--output_dir", required=True)
    p.add_argument(
        "--dump_patient_csv", default=None,
        help="2026-08-18追加: 患者単位のワイド形式データ(必要列のみ)をこのパスにCSV保存する。"
             "diamond_compare_cohorts_stats.py で外部コホートとの統計検定に使う。"
             "ローカルファイルのみに保存し、チャット等には出力しない設計。",
    )
    return p.parse_args()


def read_csv_robust(path: Path) -> pd.DataFrame:
    """ExcelのCSV書き出しはエンコーディングが環境依存になりやすいため、複数試す。"""
    last_err = None
    for enc in ["utf-8-sig", "utf-8", "cp932", "shift_jis", "mac_roman"]:
        try:
            return pd.read_csv(path, encoding=enc)
        except (UnicodeDecodeError, UnicodeError) as e:
            last_err = e
            continue
    raise RuntimeError(f"{path} を読めるエンコーディングが見つかりませんでした: {last_err}")


def normalize_colname(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(name).strip().lower())


def resolve_col(df: pd.DataFrame, wanted: str) -> Optional[str]:
    if wanted in df.columns:
        return wanted
    norm_wanted = normalize_colname(wanted)
    # ⚠️ 全角のみ(漢字・かな)の列名は正規化すると空文字列になり、
    # 別の無関係な全角オンリー列と誤って一致してしまうため、あいまい一致を無効化する。
    # (例: "主病巣側"と"患者番号"はどちらも正規化後は""になり得る)
    if not norm_wanted:
        return None
    for c in df.columns:
        if normalize_colname(c) == norm_wanted:
            return c
    return None


def find_by_keyword(df: pd.DataFrame, keyword: str) -> Optional[str]:
    for c in df.columns:
        if keyword.lower() in str(c).lower():
            return c
    return None


# ---------------------------------------------------------------------------
# 「手術術式」自由記述テキストからの構造化変数の推定（2026-08-16追加）
#
# 「手術術式」列は患者ごとにほぼユニークな自由記述だが、実データを見ると
# 一定のパターン（アプローチ／切除範囲／リンパ節郭清／左右）で書かれている
# ことが多い。既存の構造化列（Approach/肺切除範囲/切除肺葉/ﾘﾝﾊﾟ節郭清）は
# 欠測率が高い（Approach欠測17%、切除肺葉欠測31%）ため、自由記述側から
# 同等の情報を抽出できれば、より完全なTable1を作れる可能性がある。
#
# 表記揺れの例（実データで確認したもの）:
#   アプローチ: ロボット支援(下)/RATS、完全胸腔鏡下/完全鏡視下、
#               胸腔鏡補助下/胸腔鏡下（補助の有無混在）、開胸(下)、
#               cVATS/hVATS/VATS（本人指摘2026-08-16: 接頭辞なしの単独"VATS"は
#               完全鏡視下を意味する。hVATSのみ胸腔鏡補助下）
#   切除範囲:   ○葉切除（肺葉切除）、区域切除/亜区域切除/底区域切除/
#               舌区域切除/大区域切除/S+数字表記（区域切除）、
#               部分切除（wedge）、全摘（pneumonectomy）
#   郭清:       ND1a/ND1b/ND2a-1/ND2a-2等の記号、サンプリング、
#               リンパ節郭清/廓清（表記ゆれ）
#   左右:       文字列中の「右」「左」の出現
#
# 判定の優先順位は「より侵襲的/特異的な所見を優先」する設計（例:
# 「右肺上葉切除+右肺下葉部分切除」は部分切除も含むが、より侵襲的な
# 肺葉切除を主たる切除範囲として扱う）。あくまで文字列パターンに基づく
# 推定であり、100%の精度は保証しない。既存の構造化列との一致率を
# QCレポートに出力するので、乖離が大きい場合は個別に確認すること。
# ---------------------------------------------------------------------------

def classify_approach_from_text(text: object) -> Optional[str]:
    if pd.isna(text):
        return None
    t = str(text)
    if re.search(r"ロボット支援|RATS", t, re.IGNORECASE):
        return "ロボット支援下(RATS)"
    if re.search(r"完全胸腔鏡下|完全鏡視下|完全鏡視|cVATS", t, re.IGNORECASE):
        return "完全胸腔鏡下(cVATS相当)"
    if re.search(r"hVATS", t, re.IGNORECASE):
        return "胸腔鏡補助下(hVATS相当)"
    # ⚠️ 2026-08-16 本人指摘: h/c等の接頭辞なしの単独"VATS"表記は完全鏡視下を意味する。
    # \bVATS\b は使わない: Pythonの正規表現はデフォルトで漢字も\w(単語構成文字)とみなすため、
    # 「VATS右上葉切除」のように直後に漢字が続くと\bが境界と認識せずマッチに失敗する。
    # 代わりに直前がアルファベットでないことだけを否定先読みで確認する。
    if re.search(r"(?<![a-zA-Z])VATS", t, re.IGNORECASE):
        return "完全胸腔鏡下(cVATS相当)"
    if re.search(r"胸腔鏡", t):  # "胸腔鏡補助下" と 補助なしの "胸腔鏡下"（日本語表記、修飾語なし）
        return "胸腔鏡補助下(hVATS相当)"
    if re.search(r"開胸", t):
        return "開胸(thoracotomy)"
    return "不明(文字列からアプローチ判定不能)"


def classify_extent_from_text(text: object) -> Optional[str]:
    if pd.isna(text):
        return None
    t = str(text)
    if re.search(r"全摘", t):
        return "肺全摘(pneumonectomy)"
    if re.search(r"葉切除", t):  # 例: 上葉切除/下葉切除/上中葉切除/中下葉切除
        return "肺葉切除(lobectomy)"
    if re.search(r"区域切除|亜区域切除|底区域切除|舌区域切除|大区域切除|S\d[a-zA-Z0-9+]*区域|S\d[a-zA-Z0-9+]*切除", t):
        return "区域切除(segmentectomy)"
    if re.search(r"部分切除", t):
        return "部分切除(wedge)"
    return "不明(文字列から切除範囲判定不能)"


def classify_lymph_node_from_text(text: object) -> Optional[str]:
    if pd.isna(text):
        return None
    t = str(text)
    if re.search(r"ND\s*2", t, re.IGNORECASE):
        return "郭清(ND2相当)"
    if re.search(r"ND\s*1", t, re.IGNORECASE):
        return "郭清(ND1相当)"
    if re.search(r"サンプリング|sampling", t, re.IGNORECASE):
        return "サンプリングのみ"
    if re.search(r"郭清|廓清", t):
        return "郭清(詳細不明)"
    return "記載なし(文字列からは郭清情報なし)"


def classify_side_from_text(text: object) -> Optional[str]:
    if pd.isna(text):
        return None
    t = str(text)
    has_right = "右" in t
    has_left = "左" in t
    if has_right and has_left:
        return "両側/複数側言及"
    if has_right:
        return "右"
    if has_left:
        return "左"
    return "不明(文字列に左右の記載なし)"


def fix_percent_decimal_point(
    series: pd.Series,
    plausible_max: float = 300.0,
    accept_max: float = 200.0,
    id_series: Optional[pd.Series] = None,
) -> "tuple[pd.Series, pd.DataFrame]":
    """%予測値（FVC(%)/FEV1.0(%)等）の「小数点抜け」を検出・復元する。

    実データで発見（2026-08-16）: FVC(%)の最大値11706.0は117.06、
    FEV1.0(%)の最大値72373.0は72.373と、小数点を復元すると両方とも
    臨床的に妥当な範囲に収まる。0〜300%の生理学的範囲を外れる値について、
    10/100/1000で割って範囲内に収まるものがあれば「小数点抜けの復元」候補とする。

    ⚠️ 2026-08-16追加（本人指摘）: 復元候補が求まっても、それ自体がなお
    臨床的に高すぎる場合（実例: FEV1.0(%)が2780.0→278.0に「復元」できたが、
    278%はFEV1.0の%予測値としては依然として異常）は採用せず除外(NaN)する。
    accept_maxがその二段目の閾値（既定200%、plausible_max=300%より厳しい）。
    値は捏造しない：復元できない、または復元後も高すぎるものは常にNaN。

    戻り値: (補正後のSeries, 監査用DataFrame[影響を受けた行のみ])
    """
    s = pd.to_numeric(series, errors="coerce")
    fixed = s.copy()
    audit_rows = []
    out_of_range = (s < 0) | (s > plausible_max)
    for idx in s[out_of_range].index:
        raw = s.loc[idx]
        candidate_recovered = None
        for divisor in (10.0, 100.0, 1000.0):
            candidate = raw / divisor
            if 0 <= candidate <= plausible_max:
                candidate_recovered = candidate
                break
        if candidate_recovered is None:
            status = "復元不能(要手動確認)"
            final_value = np.nan
        elif candidate_recovered > accept_max:
            status = f"復元候補{candidate_recovered:.2f}も依然高値のため除外(要手動確認)"
            final_value = np.nan
        else:
            status = "小数点復元"
            final_value = candidate_recovered
        fixed.loc[idx] = final_value
        audit_rows.append({
            "row_index": idx,
            "id": id_series.loc[idx] if id_series is not None and idx in id_series.index else None,
            "raw_value": raw,
            "recovered_value": candidate_recovered,
            "accepted_value": final_value,
            "status": status,
        })
    audit_df = pd.DataFrame(audit_rows)
    return fixed, audit_df


def summarize_continuous(series: pd.Series) -> dict:
    s = pd.to_numeric(series, errors="coerce")
    n = int(s.notna().sum())
    missing = int(s.isna().sum())
    if n == 0:
        return {
            "n": 0, "missing": missing, "mean": np.nan, "sd": np.nan,
            "median": np.nan, "q1": np.nan, "q3": np.nan, "min": np.nan, "max": np.nan,
        }
    return {
        "n": n, "missing": missing,
        "mean": float(s.mean()), "sd": float(s.std()),
        "median": float(s.median()), "q1": float(s.quantile(0.25)), "q3": float(s.quantile(0.75)),
        "min": float(s.min()), "max": float(s.max()),
    }


MAX_CATEGORIES_SHOWN = 15


def summarize_categorical(series: pd.Series, max_categories: int = MAX_CATEGORIES_SHOWN) -> pd.DataFrame:
    """カテゴリ度数を集計する。カテゴリ数が多すぎる列（自由記述に近い術式名等）が
    Table1を数百行に膨らませてしまうのを防ぐため、上位max_categories件に絞り、
    残りは「その他(N件)」1行にまとめる。"""
    s = series.astype("object")
    missing = int(s.isna().sum())
    counts = s.value_counts(dropna=True)
    total = int(counts.sum())
    rows = []
    top = counts.iloc[:max_categories]
    rest = counts.iloc[max_categories:]
    for cat, n in top.items():
        rows.append({"category": cat, "n": int(n), "pct": (n / total * 100) if total else np.nan})
    if len(rest):
        rows.append({
            "category": f"その他({len(rest)}カテゴリ)",
            "n": int(rest.sum()),
            "pct": (rest.sum() / total * 100) if total else np.nan,
        })
    rows.append({"category": "(missing)", "n": missing, "pct": np.nan})
    return pd.DataFrame(rows)


def build_table1(df: pd.DataFrame, continuous_vars, categorical_vars) -> pd.DataFrame:
    records = []
    for label, col in continuous_vars:
        resolved = resolve_col(df, col)
        if resolved is None:
            records.append({"variable": label, "col": col, "type": "continuous", "note": "列が見つかりません"})
            continue
        stat = summarize_continuous(df[resolved])
        records.append({"variable": label, "col": resolved, "type": "continuous", **stat})
    for label, col in categorical_vars:
        resolved = resolve_col(df, col)
        if resolved is None:
            records.append({"variable": label, "col": col, "type": "categorical", "note": "列が見つかりません"})
            continue
        cat_df = summarize_categorical(df[resolved])
        for _, r in cat_df.iterrows():
            records.append({
                "variable": label, "col": resolved, "type": "categorical",
                "category": r["category"], "n": r["n"], "pct": r["pct"],
            })
    return pd.DataFrame(records)


def markdown_table(df: pd.DataFrame) -> str:
    try:
        return df.to_markdown(index=False)
    except ImportError:
        cols = df.columns.tolist()
        lines = ["| " + " | ".join(cols) + " |", "|" + "---|" * len(cols)]
        for _, row in df.iterrows():
            lines.append("| " + " | ".join(str(row[c]) for c in cols) + " |")
        return "\n".join(lines)


def _load_excel_sheets_openpyxl_stream(path: Path, sheet_names: List[str]) -> Dict[str, Optional[pd.DataFrame]]:
    """フォールバック: openpyxlのread_onlyストリーミングモードで手動にDataFrame化する。
    pandasのengine_kwargsに対応していない古いバージョン向け。"""
    import openpyxl

    result: Dict[str, Optional[pd.DataFrame]] = {}
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    try:
        for name in sheet_names:
            if name not in wb.sheetnames:
                result[name] = None
                continue
            ws = wb[name]
            row_iter = ws.iter_rows(values_only=True)
            try:
                header_raw = next(row_iter)
            except StopIteration:
                result[name] = pd.DataFrame()
                continue
            header = [str(h) if h is not None else f"Unnamed_{i}" for i, h in enumerate(header_raw)]
            ncol = len(header)
            data = []
            for r in row_iter:
                if len(r) < ncol:
                    r = tuple(r) + (None,) * (ncol - len(r))
                elif len(r) > ncol:
                    r = r[:ncol]
                data.append(r)
            result[name] = pd.DataFrame(data, columns=header)
    finally:
        wb.close()
    return result


def load_excel_sheets_fast(path: Path, sheet_names: List[str]) -> Dict[str, Optional[pd.DataFrame]]:
    """指定シートだけを高速に読み込む（ブックは1回だけ開く）。

    通常の pd.read_excel(path, sheet_name=...) は呼び出すたびにファイル全体
    （スタイル情報・全シート構造）を毎回パースし直すため、シート数が多く
    列数も多い巨大なxlsx（本ケースは30シート以上、幅300列超のシートも複数）
    に対して同じファイルを複数回開くと非常に遅くなる。
    まずpandasの`engine_kwargs={"read_only": True}`（スタイルを読まない
    openpyxlのストリーミングモード、pandas>=1.3で対応）を試し、
    古いpandasで使えない場合は手動ストリーミング読込にフォールバックする。
    """
    try:
        xls = pd.ExcelFile(path, engine="openpyxl", engine_kwargs={"read_only": True})
        result: Dict[str, Optional[pd.DataFrame]] = {}
        for name in sheet_names:
            if name not in xls.sheet_names:
                result[name] = None
                continue
            result[name] = pd.read_excel(xls, sheet_name=name)
        return result
    except TypeError:
        print("[INFO] pandasがengine_kwargsに未対応のため、openpyxl直接ストリーミングにフォールバックします")
        return _load_excel_sheets_openpyxl_stream(path, sheet_names)


def run_qc_checks(base: pd.DataFrame, factor: pd.DataFrame, join_key_candidates: List[str], output_dir: Path) -> None:
    """Table1を作る前の品質チェック。値そのものは出さず、集計・件数のみ報告・保存する。"""
    lines: List[str] = []
    lines.append(f"[QC] {BASE_SHEET}: {len(base)}行 / {FACTOR_SHEET}: {len(factor)}行")

    # 1. 結合キー候補ごとの重複チェック
    for key in join_key_candidates:
        if key in base.columns:
            dup_n = int(base[key].duplicated(keep=False).sum())
            lines.append(f"[QC] {BASE_SHEET}: '{key}'の重複行数={dup_n}" + ("（要確認）" if dup_n else ""))
        if key in factor.columns:
            dup_n = int(factor[key].duplicated(keep=False).sum())
            lines.append(f"[QC] {FACTOR_SHEET}: '{key}'の重複行数={dup_n}" + ("（要確認）" if dup_n else ""))

    # 2. 主要変数の欠測率
    lines.append("[QC] 主要変数の欠測率（Table1候補列のうち実在するもの）:")
    for label, col in CONTINUOUS_VARS + CATEGORICAL_VARS:
        resolved = resolve_col(base, col)
        if resolved is None:
            continue
        missing_pct = base[resolved].isna().mean() * 100
        lines.append(f"    {label}({resolved}): 欠測{missing_pct:.1f}%")

    # 3. 年齢の異常値チェック（値は出さず、件数のみ）
    age_col = resolve_col(base, "PRE-age")
    if age_col is not None:
        age_num = pd.to_numeric(base[age_col], errors="coerce")
        n_out_of_range = int(((age_num < 0) | (age_num > 110)).sum())
        n_non_numeric = int(base[age_col].notna().sum() - age_num.notna().sum())
        lines.append(f"[QC] 年齢: 0〜110の範囲外={n_out_of_range}件, 数値変換不能={n_non_numeric}件")

    # 3b. 肺機能%予測値の外れ値チェック（2026-08-16追加）
    # ⚠️ 実データで発覚: FVC(%)/FEV1.0(%)の平均・SDが中央値・IQRと比べて異常に
    # 大きく乖離していた（例: FEV1.0(%) 中央値76.7に対しSD=2386）。生理学的に
    # あり得ない範囲(<0 or >300)の値がごく少数混入しているとmean/SDが大きく
    # 歪むため、範囲外件数とmin/maxを事前に報告する（値そのものは個別に出さない）。
    for label, col in [("FVC(%)", "PRE-FVC(%)"), ("FEV1.0(%)", "PRE-FEV1.0(%)")]:
        resolved = resolve_col(factor, col)
        if resolved is None:
            continue
        vals = pd.to_numeric(factor[resolved], errors="coerce")
        if vals.notna().sum() == 0:
            continue
        n_extreme = int(((vals < 0) | (vals > 300)).sum())
        flag = "（要確認: mean/SDが信頼できない可能性）" if n_extreme else ""
        lines.append(
            f"[QC] {label}: 範囲=[{vals.min():.1f}, {vals.max():.1f}], "
            f"生理学的にあり得ない値(<0 or >300)={n_extreme}件{flag}"
        )

    # 4. カテゴリ変数の高カーディナリティ警告（自由記述に近い列がTable1を膨らませないか事前確認）
    for label, col in CATEGORICAL_VARS:
        resolved = resolve_col(base, col)
        if resolved is None:
            continue
        n_unique = int(base[resolved].nunique(dropna=True))
        if n_unique > MAX_CATEGORIES_SHOWN:
            lines.append(
                f"[QC][WARN] {label}({resolved}): カテゴリ数={n_unique}件と多く、自由記述に近い可能性。"
                f"Table1では上位{MAX_CATEGORIES_SHOWN}件+その他にまとめられます。"
            )

    report = "\n".join(lines)
    print(report)
    (output_dir / "diamond_table1_dev_qc_report.txt").write_text(report, encoding="utf-8")


def main() -> int:
    args = parse_args()
    output_dir = Path(args.output_dir).expanduser()
    output_dir.mkdir(parents=True, exist_ok=True)

    if args.base_csv and args.factor_csv:
        base_path = Path(args.base_csv).expanduser()
        factor_path = Path(args.factor_csv).expanduser()
        print(f"[INFO] CSV事前エクスポート版を読み込み中: {base_path.name} / {factor_path.name}")
        base = read_csv_robust(base_path)
        factor = read_csv_robust(factor_path)
    else:
        case_file = Path(args.case_file).expanduser()
        print(
            f"[INFO] {case_file.name} を1回だけ開いて{BASE_SHEET}/{FACTOR_SHEET}を読み込み中... "
            f"(巨大・肥大化したxlsxだと非常に遅い場合があります。遅ければCtrl+Cで止めて、"
            f"Excelで両シートをCSVエクスポートし --base_csv/--factor_csv で渡してください)"
        )
        sheets = load_excel_sheets_fast(case_file, [BASE_SHEET, FACTOR_SHEET])
        base = sheets.get(BASE_SHEET)
        factor = sheets.get(FACTOR_SHEET)
        if base is None:
            raise RuntimeError(f"シート'{BASE_SHEET}'が見つかりません。シート名を確認してください。")
        if factor is None:
            raise RuntimeError(f"シート'{FACTOR_SHEET}'が見つかりません。シート名を確認してください。")

    # ⚠️ openpyxlのread_onlyモードは、Excelファイルの「使用範囲(dimension)」が
    # 実データより過大に記録されている場合（本ケースで確認: 921行のはずが2765行、
    # 差分は全列が空白の"幽霊行"）、その幽霊行までそのまま読み込んでしまう既知の癖がある。
    # 完全に空白の行を除去して、真のデータ範囲に戻す。
    n_base_raw, n_factor_raw = len(base), len(factor)
    base = base.dropna(how="all").reset_index(drop=True)
    factor = factor.dropna(how="all").reset_index(drop=True)
    if len(base) != n_base_raw or len(factor) != n_factor_raw:
        print(
            f"[WARN] 完全に空白の行を除去しました（使用範囲が実データより過大だった可能性）: "
            f"{BASE_SHEET} {n_base_raw}→{len(base)}行, {FACTOR_SHEET} {n_factor_raw}→{len(factor)}行"
        )

    print(f"[INFO] {BASE_SHEET}: {len(base)}行, {len(base.columns)}列")
    print(f"[INFO] {FACTOR_SHEET}: {len(factor)}行, {len(factor.columns)}列")

    # Table1を作る前の品質チェック（値は出さず集計のみ、diamond_table1_dev_qc_report.txtに保存）
    run_qc_checks(base, factor, JOIN_KEY_CANDIDATES, output_dir)

    # 全列名をファイルに保存（診断用、値は含まない・短いテキストなので安全）
    dump_path = output_dir / "diamond_case_file_columns_dump.txt"
    dump_path.write_text(
        f"[{BASE_SHEET}] ({len(base.columns)}列)\n" + "\n".join(map(str, base.columns))
        + f"\n\n[{FACTOR_SHEET}] ({len(factor.columns)}列)\n" + "\n".join(map(str, factor.columns)),
        encoding="utf-8",
    )
    print(f"[OK] 全列名一覧を保存（未確認だった60列より後ろも含む）: {dump_path}")

    # キーワードで追加候補を自動検出（手術時間・出血量・術式等）
    extra_continuous = []
    for kw in KEYWORD_CANDIDATES["continuous"]:
        c = find_by_keyword(base, kw)
        if c:
            extra_continuous.append((f"{kw}(自動検出:{c})", c))
    extra_categorical = []
    for kw in KEYWORD_CANDIDATES["categorical"]:
        c = find_by_keyword(base, kw)
        if c:
            extra_categorical.append((f"{kw}(自動検出:{c})", c))
    if extra_continuous or extra_categorical:
        print(
            f"[INFO] キーワード自動検出で追加候補: "
            f"連続変数{[c for _, c in extra_continuous]}, カテゴリ変数{[c for _, c in extra_categorical]}"
        )
    else:
        print(
            "[INFO] キーワード自動検出（手術時間・出血量・術式等）は見つかりませんでした"
            "（本当の除外なし921例シートには含まれない可能性。列名一覧dumpを確認してください）"
        )

    # 結合キーの決定と身長体重等の結合
    join_key = None
    for key in JOIN_KEY_CANDIDATES:
        if key in base.columns and key in factor.columns:
            join_key = key
            break
    merged = base
    if join_key:
        wanted_factor_cols = [c for _, c in FACTOR_CONTINUOUS_VARS + FACTOR_CATEGORICAL_VARS if c in factor.columns]
        factor_cols_needed = [join_key] + wanted_factor_cols
        factor_sub = factor[factor_cols_needed].drop_duplicates(subset=[join_key])
        merged = base.merge(factor_sub, on=join_key, how="left", suffixes=("", "_factor"))
        if wanted_factor_cols:
            n_matched = int(merged[wanted_factor_cols].notna().any(axis=1).sum())
            print(f"[INFO] '{join_key}'で結合: {FACTOR_SHEET}側の情報が付与できた症例={n_matched}/{len(base)}")
    else:
        print(f"[WARN] 結合キー({JOIN_KEY_CANDIDATES})が両シートに見つからず、身長体重等は結合できませんでした")

    if "身長(cm)" in merged.columns and "体重(kg)" in merged.columns:
        h_m = pd.to_numeric(merged["身長(cm)"], errors="coerce") / 100.0
        w_kg = pd.to_numeric(merged["体重(kg)"], errors="coerce")
        merged["_bmi_computed"] = w_kg / (h_m ** 2)

    # %予測値の小数点抜けを検出・復元（2026-08-16追加）
    # 実データでFVC(%)最大値11706.0→117.06、FEV1.0(%)最大値72373.0→72.373と、
    # 小数点を復元すると臨床的に妥当な範囲に収まることが判明。復元結果はTable1に
    # 使い、監査用に影響を受けた症例（患者番号＋補正前後の値）を別ファイルに保存する。
    id_series_for_audit = merged[join_key] if join_key else None
    fixed_pct_vars = []
    for label, col in [("FVC(%)", "PRE-FVC(%)"), ("FEV1.0(%)", "PRE-FEV1.0(%)")]:
        resolved = resolve_col(merged, col)
        if resolved is None:
            continue
        fixed_series, audit_df = fix_percent_decimal_point(merged[resolved], id_series=id_series_for_audit)
        clean_col = f"_{resolved}_fixed"
        merged[clean_col] = fixed_series
        fixed_pct_vars.append((f"{label}(小数点補正後)", clean_col))
        if len(audit_df):
            audit_path = output_dir / f"diamond_table1_dev_percent_fix_audit_{resolved.replace('(', '').replace(')', '').replace('%', 'pct').replace('.', '')}.csv"
            audit_df.to_csv(audit_path, index=False)
            n_recovered = int((audit_df["status"] == "小数点復元").sum())
            n_excluded = len(audit_df) - n_recovered
            print(
                f"[INFO] {label}: 小数点抜けを{n_recovered}件復元、{n_excluded}件はNaN除外"
                f"（復元不能、または復元後もなお高値のため）"
                f"（監査用ファイル: {audit_path.name}）"
            )

    # 「手術術式」自由記述から構造化変数（アプローチ／切除範囲／郭清／左右）を推定
    proc_col = find_by_keyword(merged, "手術術式")
    derived_categorical_vars = []
    if proc_col:
        merged["_derived_approach"] = merged[proc_col].apply(classify_approach_from_text)
        merged["_derived_extent"] = merged[proc_col].apply(classify_extent_from_text)
        merged["_derived_lymph_node"] = merged[proc_col].apply(classify_lymph_node_from_text)
        merged["_derived_side"] = merged[proc_col].apply(classify_side_from_text)
        derived_categorical_vars = [
            ("アプローチ(術式文字列から推定)", "_derived_approach"),
            ("切除範囲(術式文字列から推定)", "_derived_extent"),
            ("リンパ節郭清(術式文字列から推定)", "_derived_lymph_node"),
            ("左右(術式文字列から推定)", "_derived_side"),
        ]

        n_total = int(merged[proc_col].notna().sum())
        summary_bits = []
        for label, col in [
            ("アプローチ", "_derived_approach"), ("切除範囲", "_derived_extent"),
            ("郭清", "_derived_lymph_node"), ("左右", "_derived_side"),
        ]:
            n_unknown = int(merged[col].astype(str).str.startswith(("不明", "記載なし")).sum())
            summary_bits.append(f"{label}判定不能={n_unknown}/{n_total}")
        print(f"[INFO] '{proc_col}'から構造化変数を推定（表記揺れ対応）: " + ", ".join(summary_bits))

        # 既存の構造化Approach列(INTRA-Approach)との一致率チェック（あれば、QCレポートに追記）
        approach_struct_col = find_by_keyword(merged, "INTRA-Approach")
        if approach_struct_col:
            approach_code_map = {
                "0": "胸腔鏡補助下(hVATS相当)", "0.0": "胸腔鏡補助下(hVATS相当)",
                "1": "完全胸腔鏡下(cVATS相当)", "1.0": "完全胸腔鏡下(cVATS相当)",
                "2": "ロボット支援下(RATS)", "2.0": "ロボット支援下(RATS)",
                "3": "開胸(thoracotomy)", "3.0": "開胸(thoracotomy)",
            }
            both = merged[[approach_struct_col, "_derived_approach"]].dropna()
            if len(both):
                mapped_struct = both[approach_struct_col].astype(str).map(approach_code_map)
                agree = int((mapped_struct == both["_derived_approach"]).sum())
                agree_line = (
                    f"[QC] アプローチ: 構造化列(INTRA-Approach)と文字列推定の一致="
                    f"{agree}/{len(both)}件 ({agree / len(both) * 100:.1f}%)"
                )
                print(agree_line)
                with open(output_dir / "diamond_table1_dev_qc_report.txt", "a", encoding="utf-8") as f:
                    f.write("\n" + agree_line + "\n")
    else:
        print("[INFO] '手術術式'列が見つからず、文字列からの構造化変数推定はスキップしました")

    # 2026-08-18追加: diamond_compare_cohorts_stats.py（両コホート統計検定）用に、
    # 患者単位のワイド形式データを必要列だけ選んでCSV保存する（ローカルのみ、チャット非出力）。
    if args.dump_patient_csv:
        dump_cols_wanted = [
            ("PRE-age", "age"),
            ("PRE-sex", "sex"),
            ("PRE-smoking or never", "smoking_ever"),
            ("PRE-S.I.", "si"),
            ("PRE-PS", "ps"),
            ("PRE-COPD", "copd"),
            ("PRE-FEV1.0", "fev1_ml"),
            ("PRE-FVC", "fvc_ml"),
            ("疾患の種類(1)", "primary_disease"),
            ("INTRA-Approach", "approach_struct_code"),
            ("_derived_extent", "resection_extent"),
            ("INTRA-adhesion", "adhesion"),
            ("肺瘻0or1", "pal_proxy"),
        ]
        dump_df = pd.DataFrame()
        for src_wanted, out_name in dump_cols_wanted:
            # ⚠️ 2026-08-18修正: "INTRA-Approach"は実データでは
            # "INTRA-Approach(hVATS=0, cVATS=1, RATS=2, thoracotomy=3)"のように
            # 注記付きの列名になっており、resolve_col(完全一致/正規化一致のみ)では
            # 見つからない。find_by_keyword(部分一致)で探す必要がある
            # （QCの一致率チェックで既に使っているapproach_struct_colと同じロジック）。
            if src_wanted == "INTRA-Approach":
                resolved = find_by_keyword(merged, "INTRA-Approach")
            else:
                resolved = resolve_col(merged, src_wanted)
            if resolved is not None:
                dump_df[out_name] = merged[resolved]
            else:
                dump_df[out_name] = np.nan
                print(f"[WARN] --dump_patient_csv: 列'{src_wanted}'が見つからず{out_name}は全欠損になります")
        dump_path = Path(args.dump_patient_csv).expanduser()
        dump_path.parent.mkdir(parents=True, exist_ok=True)
        dump_df.to_csv(dump_path, index=False)
        print(f"[OK] 患者単位ワイドデータ(n={len(dump_df)})を保存: {dump_path}")

    # ⚠️ FVC(%)/FEV1.0(%)は生の値(小数点抜けを含み得る)と補正後の値を両方Table1に載せる。
    # 生の値を消さないのは透明性のため（何を・どう補正したかを常に追跡できるように）。
    continuous_vars = CONTINUOUS_VARS + extra_continuous + FACTOR_CONTINUOUS_VARS + fixed_pct_vars
    if "_bmi_computed" in merged.columns:
        continuous_vars = continuous_vars + [("BMI(算出値)", "_bmi_computed")]
    categorical_vars = CATEGORICAL_VARS + extra_categorical + FACTOR_CATEGORICAL_VARS + derived_categorical_vars

    table1 = build_table1(merged, continuous_vars, categorical_vars)
    csv_path = output_dir / "table1_dev_cohort.csv"
    md_path = output_dir / "table1_dev_cohort.md"
    table1.to_csv(csv_path, index=False)

    n_base = len(base)
    ref_ns = list(MODEL_COHORT_N_REFERENCE.values())
    missing_lo = min(ref_ns) - n_base
    missing_hi = max(ref_ns) - n_base
    header_note = (
        f"# 単施設(DIAMOND)コホート Table1\n\n"
        f"母集団: `{BASE_SHEET}`シート, n={n_base}\n\n"
        f"**重要な前提（2026-08-16本人確認）**: この{n_base}例は、実際に時系列モデルの学習に"
        f"使われているコホート（参考値: {MODEL_COHORT_N_REFERENCE}）の部分集合であり、"
        f"臨床背景データが揃っている症例のみです（欠測 約{missing_lo}〜{missing_hi}例）。"
        f"論文ではモデル学習コホート全体（n≈1020〜1025）とTable1の母数（n={n_base}）が"
        f"異なる点、および欠測症例の特性未検討である点をMethods/Limitationsに明記してください。\n\n"
    )
    md_path.write_text(header_note + markdown_table(table1), encoding="utf-8")

    print(f"[OK] Table1を保存: {csv_path} / {md_path}")
    print(
        f"[NOTE] このTable1の母数(n={n_base})は、モデル学習コホート"
        f"(参考: {MODEL_COHORT_N_REFERENCE})の部分集合（背景データ完備例）です。"
        f"欠測症例の特性（ランダム欠測か系統的な違いがあるか）は未検討のため、"
        f"可能であれば欠測群との比較も検討してください。"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
